"""
File upload analysis — lets the chatbot see images and read documents.

Images (png/jpg/webp/gif) are sent to Gemini's vision endpoint as base64 data
URIs; documents (PDF/DOCX/XLSX/TXT/CSV/JSON/MD) get their text extracted in
code and included in the prompt. The answer is grounded ONLY in the uploaded
content (plus learned business rules for context).

Used by app.py /ask when the request is multipart/form-data with files.
"""
from __future__ import annotations

import base64
import io
import os

from dotenv import load_dotenv
from openai import OpenAI

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

_client = OpenAI(
    api_key=os.getenv("GEMINI_API_KEY"),
    base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
)
_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

MAX_FILES = 5
MAX_FILE_BYTES = 10 * 1024 * 1024   # 10 MB per file
_MAX_DOC_CHARS = 50000              # per document, keeps the prompt sane

_IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "webp": "image/webp", "gif": "image/gif",
}
_TEXT_EXT = {"txt", "csv", "json", "md", "log", "sql", "tsv"}

ALLOWED_EXTENSIONS = set(_IMAGE_MIME) | _TEXT_EXT | {"pdf", "docx", "xlsx"}


def _ext(filename: str) -> str:
    return (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()


# =========================================================
# DOCUMENT TEXT EXTRACTION
# =========================================================
def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        pages.append(f"--- page {i + 1} ---\n{(page.extract_text() or '').strip()}")
        if sum(len(p) for p in pages) > _MAX_DOC_CHARS:
            pages.append(f"...(stopped at page {i + 1} of {len(reader.pages)})")
            break
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    total = 0
    for ws in wb.worksheets:
        parts.append(f"--- sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            line = " | ".join("" if v is None else str(v) for v in row)
            parts.append(line)
            total += len(line)
            if total > _MAX_DOC_CHARS:
                parts.append("...(truncated)")
                wb.close()
                return "\n".join(parts)
    wb.close()
    return "\n".join(parts)


def _extract_document_text(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    try:
        if ext == "pdf":
            text = _extract_pdf(data)
        elif ext == "docx":
            text = _extract_docx(data)
        elif ext == "xlsx":
            text = _extract_xlsx(data)
        elif ext in _TEXT_EXT:
            text = data.decode("utf-8", errors="replace")
        else:
            return f"[Unsupported file type: .{ext}]"
    except Exception as e:
        return f"[Could not read this file: {e}]"
    text = text.strip()
    if not text:
        return "[The file was read but contained no extractable text — it may be a scanned image inside a PDF.]"
    if len(text) > _MAX_DOC_CHARS:
        text = text[:_MAX_DOC_CHARS] + "\n...(truncated)"
    return text


# =========================================================
# ANALYSIS
# =========================================================
def _learned_rules() -> str:
    try:
        from learning_store import lessons_block
        return lessons_block()
    except Exception:
        return ""


def _system_prompt() -> str:
    learned = _learned_rules()
    return f"""You are the WEBXPAY analytics assistant. The user has uploaded one or more files
(images and/or documents) and wants you to analyze and explain them.

RULES:
- Ground EVERYTHING in the uploaded content. Read charts, tables, screenshots and
  documents carefully and quote the actual values you see. NEVER invent numbers,
  names or facts that are not visible in the files.
- If an image is unclear or a document section is unreadable, say so plainly.
- Explain in clear, conversational language what the file contains, then answer
  the user's specific question about it. Lead with the direct answer.
- Use markdown; short tables are fine for extracted figures; format LKR amounts
  with thousand separators.
- WEBXPAY context: payment gateway company in Sri Lanka; IPG = online gateway,
  POS = physical card machines, GMV = gross merchandise value, MDR = merchant
  discount rate, RM = relationship manager.
{(chr(10) + learned + chr(10)) if learned else ""}"""


def _history_block(history) -> str:
    if not isinstance(history, list) or not history:
        return ""
    lines = []
    for m in history[-4:]:
        role = "User" if m.get("role") == "user" else "Assistant"
        lines.append(f"{role}: {str(m.get('content', ''))[:400]}")
    return "Recent conversation (context only):\n" + "\n".join(lines) + "\n\n"


def analyze_files(question: str, files: list[tuple[str, bytes]], history=None) -> dict:
    """files = [(filename, raw_bytes), ...]. Returns a payload dict shaped like
    handle_user_question's output so the /ask response path is unchanged."""
    question = (question or "").strip() or "Analyze the attached file(s) and explain what they contain."

    content: list[dict] = []
    doc_blocks = []
    names = []

    for filename, data in files[:MAX_FILES]:
        names.append(filename)
        ext = _ext(filename)
        if ext in _IMAGE_MIME:
            b64 = base64.b64encode(data).decode("ascii")
            content.append({
                "type": "image_url",
                "image_url": {"url": f"data:{_IMAGE_MIME[ext]};base64,{b64}"},
            })
        else:
            doc_blocks.append(
                f"===== FILE: {filename} =====\n{_extract_document_text(filename, data)}"
            )

    text_prompt = _history_block(history)
    if doc_blocks:
        text_prompt += "UPLOADED DOCUMENT CONTENT:\n" + "\n\n".join(doc_blocks) + "\n\n"
    text_prompt += f"User question about the uploaded file(s) ({', '.join(names)}): {question}"
    content.insert(0, {"type": "text", "text": text_prompt})

    resp = _client.chat.completions.create(
        model=_MODEL,
        temperature=0.2,
        max_tokens=3000,
        reasoning_effort="none",  # disable Gemini thinking; else it eats the token budget
        messages=[
            {"role": "system", "content": _system_prompt()},
            {"role": "user", "content": content},
        ],
    )
    answer = (resp.choices[0].message.content or "").strip()
    if not answer:
        answer = "I couldn't extract anything useful from the uploaded file(s). Please try a clearer image or a text-based document."

    return {
        "question": question,
        "sql": None,
        "raw_result": [],
        "answer": answer,
        "insights": answer,
        "response_type": "file_analysis",
        "files": names,
    }
